"""Pro stage 1 (backend/api_pro.py, backend/report.py): free preview, Pro code redemption, credits taken atomically and only
for a built report, and the report's content rules. Temp DB with the simulated dataset, real HTTP server."""
import hashlib
import io
import json
import os
import re
import tempfile
import threading
import unittest
import urllib.error
import urllib.request

_tmp = tempfile.mkdtemp()
os.environ.setdefault("AGELENS_DB", os.path.join(_tmp, "t.db"))

from http.server import ThreadingHTTPServer  # noqa: E402
from backend import api_pro as P, db, server  # noqa: E402

CODES = {"five": "PRO-FIVE-5555", "one": "PRO-ONEX-1111", "old": "PRO-OLDX-0000"}
H = lambda c: hashlib.sha256(c.encode()).hexdigest()
ENTRIES = [{"id": "pro_5", "label": "Teacher", "credits": 5, "expires": "2099-12-31", "sha256": H(CODES["five"])},
           {"id": "pro_1", "label": "Tester", "credits": 1, "expires": "2099-12-31", "sha256": H(CODES["one"])},
           {"id": "pro_o", "label": "Old", "credits": 5, "expires": "2020-01-01", "sha256": H(CODES["old"])}]
ASK = {"question": "I'm 38 with eight years in procurement. What should I clarify?",
       "profile": {"age": 38, "occupation": "Procurement", "years_experience": 8, "region": "上海"}}
NOTHING = {"question": "x", "profile": {"occupation": "Astronaut", "region": "上海"}, "filters": {"q": "zzzz-no-such-ad"}}
BAD = ("you will be hired", "you will be rejected", "you are eligible", "you qualify", "illegal", "discriminat", "probability", "guarantee")


def seed(path):
    """The simulated dataset into this module's own DB (other test modules share the default temp DB with other ads)."""
    import subprocess, sys
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    subprocess.run([sys.executable, os.path.join(root, "scripts", "import_csv.py"), "--source", "sim"], check=True, cwd=root,
                   env=dict(os.environ, AGELENS_DB=str(path)), capture_output=True)


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from pathlib import Path
        cls._db = db.DB_PATH
        db.DB_PATH = Path(tempfile.mkdtemp()) / "pro.db"  # own DB with exactly the simulated dataset
        db.init_schema()
        server.load_routes()
        seed(db.DB_PATH)
        cls.httpd = ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
        cls.base = f"http://127.0.0.1:{cls.httpd.server_address[1]}"
        threading.Thread(target=cls.httpd.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.httpd.shutdown()
        db.DB_PATH = cls._db

    def setUp(self):
        self._old = os.environ.get("AGELENS_PRO_CODES")
        os.environ["AGELENS_PRO_CODES"] = json.dumps(ENTRIES)
        for lim in (P.fail_by_client, P.report_limiter):
            lim.reset()
        with db.connect() as c:
            c.execute("DELETE FROM pro_usage"); c.execute("DELETE FROM pro_reports")

    def tearDown(self):
        if self._old is None:
            os.environ.pop("AGELENS_PRO_CODES", None)
        else:
            os.environ["AGELENS_PRO_CODES"] = self._old

    def call(self, method, path, body=None, headers=None, raw=False):
        data = json.dumps(body).encode() if body is not None else None
        r = urllib.request.Request(self.base + path, data=data, method=method, headers=dict({"Content-Type": "application/json"}, **(headers or {})))
        try:
            with urllib.request.urlopen(r, timeout=60) as resp:
                b = resp.read()
                return resp.status, (b if raw else json.loads(b)), resp.headers
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read()), e.headers

    def pro(self, which="five"):
        s, b, _ = self.call("POST", "/api/pro/redeem", {"code": CODES[which]})
        self.assertEqual(s, 200, b)
        return {"X-Pro-Token": b["token"]}, b


class TestFreeAndGate(Base):
    def test_preview_is_free(self):
        s, o, _ = self.call("POST", "/api/report/preview", ASK)
        self.assertEqual(s, 200, o)
        self.assertEqual(len(o["sections"]), 10)
        self.assertEqual(o["summary"]["analysed"], o["summary"]["matched"])
        self.assertEqual(o["data_label"], "Demo · Synthetic data")

    def test_report_needs_pro(self):
        s, b, _ = self.call("POST", "/api/report", ASK)
        self.assertEqual(s, 402)
        self.assertTrue(b["pro_required"])
        self.assertEqual(self.call("POST", "/api/report", ASK, {"X-Pro-Token": "made-up"})[0], 402)

    def test_free_features_unaffected(self):
        self.assertEqual(self.call("POST", "/api/brief", ASK)[0], 200)
        self.assertEqual(self.call("GET", "/api/overview")[0], 200)


class TestCodes(Base):
    def test_redeem_and_status(self):
        hdr, b = self.pro()
        self.assertEqual((b["label"], b["credits_total"], b["credits_left"]), ("Teacher", 5, 5))
        self.assertEqual(self.call("GET", "/api/pro/status", headers=hdr)[1]["credits_left"], 5)
        self.assertEqual(self.call("GET", "/api/pro/status")[1], {"active": False})

    def test_bad_and_expired_codes_look_the_same(self):
        a = self.call("POST", "/api/pro/redeem", {"code": "PRO-NOPE-NOPE"})[1]
        b = self.call("POST", "/api/pro/redeem", {"code": CODES["old"]})[1]
        self.assertEqual(a, b)
        self.assertEqual(a["error"], P.BAD_CODE)

    def test_wrong_codes_limited(self):
        codes = [self.call("POST", "/api/pro/redeem", {"code": "PRO-WRON-G%03d" % i}, {"X-Forwarded-For": "10.7.7.7"})[0] for i in range(11)]
        self.assertEqual(codes[-1], 429)

    def test_revoked_code_stops_working(self):
        hdr, _ = self.pro()
        os.environ["AGELENS_PRO_CODES"] = json.dumps([e for e in ENTRIES if e["id"] != "pro_5"])
        self.assertEqual(self.call("GET", "/api/pro/status", headers=hdr)[1], {"active": False})
        self.assertEqual(self.call("POST", "/api/report", ASK, hdr)[0], 402)

    def test_signout(self):
        hdr, _ = self.pro()
        self.call("POST", "/api/pro/signout", {}, hdr)
        self.assertEqual(self.call("GET", "/api/pro/status", headers=hdr)[1], {"active": False})

    def test_only_hashes_stored(self):
        hdr, b = self.pro()
        with db.connect() as c:
            dump = json.dumps([dict(r) for r in c.execute("SELECT * FROM pro_sessions")])
        self.assertNotIn(b["token"], dump)
        self.assertNotIn(CODES["five"], dump)


class TestCredits(Base):
    def test_report_uses_one_credit_and_is_a_docx(self):
        hdr, _ = self.pro()
        s, data, hd = self.call("POST", "/api/report", ASK, hdr, raw=True)
        self.assertEqual(s, 200)
        self.assertIn("wordprocessingml", hd["Content-Type"])
        self.assertIn("attachment", hd["Content-Disposition"])
        self.assertEqual(hd["X-Credits-Left"], "4")
        self.assertEqual(self.call("GET", "/api/pro/status", headers=hdr)[1]["credits_left"], 4)
        from docx import Document
        Document(io.BytesIO(data))  # opens as Word

    def test_credits_shared_across_browsers_and_run_out(self):
        a, _ = self.pro("one")
        b, _ = self.pro("one")  # same code, another browser: same wallet
        self.assertEqual(self.call("POST", "/api/report", ASK, a, raw=True)[0], 200)
        s, body, _ = self.call("POST", "/api/report", ASK, b)
        self.assertEqual((s, body["credits_left"]), (403, 0))

    def test_empty_report_costs_nothing(self):
        hdr, _ = self.pro("one")
        s, b, _ = self.call("POST", "/api/report", NOTHING, hdr)
        self.assertEqual(s, 422, b)
        self.assertEqual(self.call("GET", "/api/pro/status", headers=hdr)[1]["credits_left"], 1)

    def test_last_credit_cannot_be_spent_twice(self):
        a, _ = self.pro("one")
        out = []
        ts = [threading.Thread(target=lambda: out.append(self.call("POST", "/api/report", ASK, a, raw=True)[0])) for _ in range(4)]
        [t.start() for t in ts]; [t.join() for t in ts]
        self.assertEqual(sorted(out), [200, 403, 403, 403])
        with db.connect() as c:
            self.assertEqual(c.execute("SELECT used FROM pro_usage WHERE code_id='pro_1'").fetchone()[0], 1)


class TestReportContent(Base):
    def text(self):
        from backend import report
        from docx import Document
        b, ctx = report.build(ASK)
        d = Document(io.BytesIO(report.to_docx(b, ctx)))
        return "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells), b, ctx

    def test_sections_labels_and_quotes(self):
        txt, b, ctx = self.text()
        for t in ("1. Summary in numbers", "9. Evidence appendix", "10. Method and limits", "Draft message"):
            self.assertIn(t, txt)
        self.assertIn("Demo · Synthetic data", txt)
        self.assertIn("simulated demonstration data", txt)
        self.assertIn("No AI model was used", txt)
        for a in ctx["ads"]:  # every ad's age wording is quoted verbatim
            for t in a.get("age_terms") or []:
                self.assertIn("“%s”" % t["snippet"], txt)

    def test_no_outcome_or_legal_claims(self):
        txt, _, _ = self.text()
        low = txt.lower().replace("does not say whether you are eligible", "")
        for w in BAD:
            self.assertNotIn(w, low, w)

    def test_deterministic(self):
        from backend import report
        a, ca = report.build(ASK)
        b, cb = report.build(ASK)
        self.assertEqual(report.outline(a, ca), report.outline(b, cb))


if __name__ == "__main__":
    unittest.main()
