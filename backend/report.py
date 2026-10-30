"""AgeLens full report (Pro): a Word document built from the same rules-only brief the Job Seeker page shows, plus context
tables and an evidence appendix. No model is called (allow_model=False), so the same request gives the same report.

Sections: cover (question, background, data label, disclaimer) · 1 summary in numbers · 2-6 the five answer sections with
the ads' wording glossed in English · 7 how your region and occupation compare (simulated dataset) · 8 questions to send, with an
editable draft message · 9 evidence appendix (every ad, its age wording as an English gloss) · 10 method and limits.
It never says whether someone qualifies, never predicts an outcome and never calls an employer's practice illegal.
outline(b) is the free preview: section titles plus the summary numbers."""
import io
from collections import Counter
from datetime import datetime, timezone

from . import brief, gloss

OCC_EN = gloss.OCC_EN        # one table for the whole app (backend/gloss.py)
CLASS_EN = gloss.CLASS_EN
_en = brief._en              # last guard: no Chinese characters in the Word file
REPORT_VERSION = "report-v1"
TITLES = ["Summary in numbers", "What the ads explicitly say", "Exceptions in the wording", "How this relates to your question",
          "What cannot be determined yet", "What you could ask next", "How your region and occupation compare",
          "Questions to send, and a draft message", "Evidence appendix", "Method and limits"]


def occ_en(v):
    return gloss.occupation_en(v) or "Not stated"


def build(body):
    """-> (brief dict, context dict). Raises ValueError when nothing matched (no report is made, so no credit is used)."""
    from . import evidence as ev
    b = brief.build_brief(body, allow_model=False)
    if not b["scope"].get("ads_analysed"):
        raise ValueError("No ads matched this question and background, so there is nothing to report. Broaden the search first.")
    prof = b["question_background"]["profile"]
    ctx = {"overview": ev.overview(), "provinces": ev.by_province(), "occupations": ev.by_occupation(limit=50), "region": prof.get("region") or "",
           "occupation_key": (b["scope"].get("filters_applied") or {}).get("occupation") or ""}
    ids = list(b["evidence_index"].keys())
    ctx["ads"] = [a for a in (ev.get_ad(i) for i in ids) if a]
    return b, ctx


def summary(b, ctx):
    """Numbers shared by the preview and the report."""
    ads = ctx["ads"]
    cls = Counter((a.get("age_class") or "none") for a in ads)
    maxes = [t.get("max_age") for a in ads for t in (a.get("age_terms") or []) if t.get("max_age") is not None]
    age = b["question_background"]["profile"].get("age")
    below = sum(1 for a in ads if any(t.get("max_age") is not None and age is not None and t["max_age"] < age for t in a.get("age_terms") or []))
    return {"matched": b["scope"]["matched_total"], "analysed": len(ads), "by_class": {CLASS_EN[k]: cls.get(k, 0) for k in CLASS_EN},
            "stated_max_ages": dict(sorted(Counter(maxes).items())), "your_age": age, "ads_stating_a_max_below_your_age": below if age is not None else None,
            "data_label": b["data_label"], "data_version": b["scope"].get("data_version")}


def outline(b, ctx):
    return {"version": REPORT_VERSION, "sections": TITLES, "summary": summary(b, ctx), "data_label": b["data_label"],
            "pages_estimate": max(4, 3 + len(ctx["ads"]) // 3)}


def _region_row(ctx):
    rows = {p["province"]: p for p in ctx["provinces"]}
    tot = ctx["overview"]
    mine = rows.get(ctx["region"]) if ctx["region"] else None
    return mine, tot


def to_docx(b, ctx):
    """-> bytes of the .docx report."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    muted = RGBColor(0x6B, 0x5B, 0x4B)

    def note(text):
        p = doc.add_paragraph(); r = p.add_run(text); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = muted
        return p

    def table(head, rows, caption=None):
        if caption:
            c = doc.add_paragraph(); r = c.add_run(caption); r.bold = True; r.font.size = Pt(10)
        t = doc.add_table(rows=1, cols=len(head)); t.style = "Light Grid Accent 1"
        for i, hd in enumerate(head):
            t.rows[0].cells[i].text = str(hd)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = "" if v is None else str(v)
        doc.add_paragraph()  # keeps consecutive tables apart
        return t

    def stmts(key):
        items = b["answer"].get(key) or []
        if not items:
            doc.add_paragraph("Nothing to report here.")
            return
        for s in items:
            p = doc.add_paragraph(style="List Bullet")
            tag = brief.TAG.get(s.get("type"), s.get("type", ""))
            p.add_run("[%s] " % tag).bold = True
            p.add_run(_en(s.get("text_en") or s["text"]))
            if s.get("quote"):
                q = doc.add_paragraph(); q.paragraph_format.left_indent = Pt(18)
                r = q.add_run("Ad wording, in English: “%s”" % _en(s.get("quote_en") or gloss.quote_en(s["quote"]))); r.font.size = Pt(10)
            if s.get("evidence_ids"):
                note("Evidence: " + ", ".join(s["evidence_ids"][:6]) + (" and %d more" % (len(s["evidence_ids"]) - 6) if len(s["evidence_ids"]) > 6 else ""))

    sm = summary(b, ctx)
    prof = b["question_background"]["profile"]
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    # cover
    doc.add_heading("AgeLens report: age wording in job ads", 0)
    p = doc.add_paragraph(); r = p.add_run("%s · rules only, no AI model · %s · %s" % (_en(b["data_label"]), REPORT_VERSION, now)); r.bold = True
    doc.add_paragraph(b["disclaimer"])
    if b.get("using_demo"):
        p = doc.add_paragraph(); r = p.add_run("All ads in this report are simulated demonstration data, not real postings."); r.bold = True
    doc.add_heading("Your question", 2); doc.add_paragraph(_en(b["question_background"]["question"]) or "(no question given)")
    doc.add_heading("Background you entered", 2)
    shown = [("Age", prof.get("age")), ("Occupation", _en(prof.get("occupation"))), ("Years of experience", prof.get("years_experience")),
             ("Region", _en(brief._region_en(prof.get("region"))) if prof.get("region") else None), ("Education", _en(prof.get("education"))),
             ("Skills / experience", _en(", ".join(prof.get("skills") or [])))]
    table(["Field", "Value"], [(k, v) for k, v in shown if v not in (None, "", [])])
    note("Your age is used only to explain conditions. No ads are hidden or ranked by it.")

    # 1 summary
    doc.add_heading("1. " + TITLES[0], 1)
    doc.add_paragraph("%d ads matched your question and background; %d were analysed." % (sm["matched"], sm["analysed"]))
    table(["Kind of age wording", "Ads"], [(k, v) for k, v in sm["by_class"].items() if v], "How the ads word age")
    if sm["stated_max_ages"]:
        table(["Stated maximum age", "Ads"], list(sm["stated_max_ages"].items()), "Maximum ages the ads state")
    if sm["ads_stating_a_max_below_your_age"] is not None:
        doc.add_paragraph("%d of these ads state a maximum age below the age you entered (%s). This describes the wording only: it does not say "
                          "whether an employer would consider you, and exceptions or other criteria may apply." % (sm["ads_stating_a_max_below_your_age"], sm["your_age"]))

    # 2-6 answer sections
    for i, hd in enumerate(b["headings"]):
        doc.add_heading("%d. %s" % (i + 2, TITLES[i + 1]), 1)
        stmts(hd["key"])

    # 7 context
    doc.add_heading("7. " + TITLES[6], 1)
    note("Shares of ads in the whole simulated dataset (%s). They describe this dataset only, not hiring or unemployment." % _en(sm["data_version"] or b["data_label"]))
    mine, tot = _region_row(ctx)
    rows = [("All regions", tot["denominator"], "%.1f%%" % (100 * tot["explicit_share"]))]
    if mine:
        rows.insert(0, (_en(brief._region_en(ctx["region"])), mine["ads"], "%.1f%%" % (100 * mine["explicit_share"])))
    occ = next((o for o in ctx["occupations"] if o["occupation"] == ctx["occupation_key"]), None)
    if occ:
        rows.insert(1 if mine else 0, (occ_en(occ["occupation"]) + " (all regions)", occ["ads"], "%.1f%%" % (100 * occ["explicit_share"])))
    table(["Group", "Ads", "Share with an explicit age condition"], rows, "Your region and occupation against the whole dataset")
    top = sorted(ctx["provinces"], key=lambda p: -p["ads"])[:8]
    num = lambda v: "—" if v is None else ("%d" % v if float(v).is_integer() else "%.1f" % v)
    table(["Region", "Ads", "Explicit age condition", "Median stated maximum"],
          [(_en(brief._region_en(p["province"])), p["ads"], "%.1f%%" % (100 * p["explicit_share"]), num(p["median_max_age"])) for p in top], "The 8 largest regions in the dataset")

    # 8 questions + draft
    doc.add_heading("8. " + TITLES[7], 1)
    qs = [_en(q.get("text_en") or q["text"]) for q in b["questions"]]
    for q in qs:
        doc.add_paragraph(q, style="List Number")
    doc.add_heading("Draft message (edit before sending)", 2)
    bg = []
    if prof.get("age") is not None:
        bg.append("I am %s" % prof["age"])
    if prof.get("years_experience") is not None:
        bg.append("with %s years of experience%s" % (prof["years_experience"], " in " + _en(prof["occupation"]) if prof.get("occupation") else ""))
    lines = ["Hello,", "", "I am interested in the position \"[job title]\" (ad [ad id]). " + (" ".join(bg) + "." if bg else ""),
             "The ad says: \"[quote the age-related wording from the ad]\".",
             "Before I apply, could you please clarify:"] + ["%d. %s" % (i + 1, q) for i, q in enumerate(qs[:3])] + ["", "Thank you for your time."]
    for ln in lines:
        doc.add_paragraph(ln)
    note("A starting point only. Replace the bracketed parts with the ad's details, keep the ad's own wording, and do not add claims about the employer.")

    # 9 appendix
    doc.add_heading("9. " + TITLES[8], 1)
    for a in ctx["ads"]:
        doc.add_heading("%s · %s" % (a["id"], _en(brief._title_en(a))), 3)
        if a.get("company_en"):
            note(_en(a["company_en"]))
        meta = [gloss.province_en(a.get("province")) or "", a.get("city_en") or gloss.city_en(a.get("city")) or "", occ_en(a.get("occupation")),
                a.get("posted_date") or "", CLASS_EN.get(a.get("age_class") or "none", "")]
        meta = [_en(x) for x in meta]
        note(" · ".join(x for x in meta if x) + (" · simulated" if a.get("is_demo") else ""))
        terms = a.get("age_terms") or []
        if not terms:
            doc.add_paragraph("No explicit age condition detected in this text (this does not mean there is no age preference).")
        if a.get("age_sentences_en"):
            note("Age wording in the English version of this simulated ad: " + _en(" / ".join(a["age_sentences_en"])))
        for t in terms:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(_en(t.get("gloss_en") or gloss.term_en(t)))
            ages = ("max %s" % t["max_age"]) if t.get("max_age") is not None else ""
            ages += ((", " if ages else "") + "min %s" % t["min_age"]) if t.get("min_age") is not None else ""
            exc = t.get("exception_en") or (gloss.exception_en(t["exception_text"]) if t.get("exception_text") else "")
            extra = [x for x in (ages, CLASS_EN.get(t.get("age_class"), ""), ("exception: " + _en(exc)) if exc else "", t.get("flag") or "") if x]
            if extra:
                p.add_run("  (" + "; ".join(extra) + ")").italic = True

    # 10 method
    doc.add_heading("10. " + TITLES[9], 1)
    for lim in b["limitations"]:
        doc.add_paragraph(lim, style="List Bullet")
    doc.add_paragraph("Method: AgeLens reads the age wording in each ad with fixed rules (%s), groups ads by the kind of wording, and shows an "
                      "English gloss of the original wording next to every statement. No AI model was used for this report. Search scope: %s." %
                      (", ".join(b["scope"].get("rule_versions") or []), _en((b["scope"].get("filters_applied") or {}).get("scope_label", ""))))
    doc.add_paragraph("Language: this report is in English only. Ad wording is shown as an English gloss of what the rules read from the "
                      "original Chinese text; the original text is kept in each ad's record in AgeLens.")
    f = doc.add_paragraph(); f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = f.add_run("AgeLens · Northwestern SEPOL 425 course project · not legal advice · no hiring decisions"); r.font.size = Pt(8); r.font.color.rgb = muted
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()
